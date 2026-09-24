import io
from datetime import date

from docx import Document
from fastapi.testclient import TestClient

from manu import drafting
from manu.api import create_app
from manu.case_state.store import CaseStore
from manu.connectors import ConnectorLadder
from manu.demo import seed

TODAY = date(2026, 9, 24)


def cases():
    store = CaseStore()
    seed(store, today=TODAY)
    return store, {c.cnr: c for c in store.all()}


def text(draft):
    return "\n".join(b.text for b in draft.blocks)


def test_a_draft_pauses_for_what_only_the_advocate_knows():
    _, by = cases()
    draft = drafting.build("adjournment", by["DLSE030005552023"], {}, as_of=TODAY)
    assert not draft.ready and set(draft.missing) == {"applicant", "reason", "previous"}
    assert "[the reason for the adjournment]" in text(draft)
    done = drafting.build(
        "adjournment",
        by["DLSE030005552023"],
        {"applicant": "defendant", "reason": "the certified copy of the sale deed is awaited from the Sub-Registrar", "previous": "3"},
        as_of=TODAY,
    )
    assert done.ready
    assert "ORDER XVII RULE 1" in text(done)
    assert "certified copy of the sale deed is awaited" in text(done)
    assert any("three adjournments" in f for f in done.flags)
    listed = next(b for b in done.blocks if "is listed before" in b.text)
    assert listed.sources == ["court record"]


def test_bail_application_uses_the_record_and_the_law_engine_with_sources():
    _, by = cases()
    aamir = by["DLSE010001232024"]
    draft = drafting.build("bail", aamir, {"provision": "483", "previous_bail": "no other bail application has been filed"}, as_of=TODAY)
    body = text(draft)
    assert draft.ready
    assert "SECTION 483" in body
    assert "Sec. 303(2) BNS" in body and "361 days" in body
    assert "no previous involvement" in body
    assert "one-third of the maximum sentence on 28.09.2026" in body
    assert any("Section 479 working" in s for b in draft.blocks for s in b.sources)
    assert any("not yet reviewed" in f for f in draft.flags)


def test_bail_asks_which_accused_when_there_are_two_and_is_not_offered_in_a_civil_suit():
    _, by = cases()
    imran = by["DLSE010007892026"]
    two = imran.model_copy(update={"accused": [*imran.accused, imran.accused[0].model_copy(update={"name": "Salim"})]})
    draft = drafting.build("bail", two, {}, as_of=TODAY)
    assert "accused" in draft.missing
    assert "[the applicant]" in text(draft)
    assert [t["key"] for t in drafting.available(by["DLSE030005552023"])] == ["adjournment"]


def test_docx_is_offered_only_when_the_draft_is_complete():
    store, by = cases()
    c = TestClient(create_app(store, ConnectorLadder([])))
    case_id = by["DLSE010001232024"].id
    assert c.post(f"/api/cases/{case_id}/drafts/bail/docx", json={"inputs": {}}).status_code == 400
    ok = c.post(
        f"/api/cases/{case_id}/drafts/bail/docx",
        json={"inputs": {"provision": "483", "previous_bail": "none has been filed"}, "on": TODAY.isoformat()},
    )
    assert ok.status_code == 200
    doc = Document(io.BytesIO(ok.content))
    assert any("REGULAR BAIL" in p.text for p in doc.paragraphs)
    assert "Advocate review required" in doc.paragraphs[-1].text
