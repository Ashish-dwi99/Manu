import io
from datetime import date

import pytest
from fastapi.testclient import TestClient

from manu import list_of_dates
from manu.api import create_app
from manu.case_state.store import CaseStore
from manu.connectors import ConnectorLadder
from manu.demo import seed
from manu.documents import DocumentError, DocumentStore, extract_pages

TODAY = date(2026, 9, 24)


def seeded():
    store = CaseStore()
    seed(store, today=TODAY)
    return store, store.by_cnr("DLSE010001232024")


def test_text_pages_split_on_form_feed():
    assert extract_pages("a.txt", b"one\fTwo\f") == ["one", "Two", ""]


def test_docx_is_read():
    from docx import Document

    doc = Document()
    doc.add_paragraph("Reply to the bail application")
    doc.add_paragraph("The applicant is not a first-time offender.")
    buffer = io.BytesIO()
    doc.save(buffer)
    pages = extract_pages("reply.docx", buffer.getvalue())
    assert "not a first-time offender" in pages[0]


def test_unsupported_and_empty_case_are_refused():
    store, case = seeded()
    documents = DocumentStore(store)
    with pytest.raises(DocumentError):
        documents.add(case.id, "photo.jpg", b"x")
    with pytest.raises(DocumentError):
        documents.add("case_missing", "a.txt", b"x")


def test_search_cites_file_and_page_and_includes_orders():
    store, case = seeded()
    documents = DocumentStore(store)
    hit = documents.search(case.id, "first-time offender")[0]
    assert (hit["name"], hit["page"]) == ("Bail-application.txt", 2)
    assert "first-time offender" in hit["quote"]
    kinds = {m["kind"] for m in documents.search(case.id, "previous involvement")}
    assert kinds == {"document", "order"}


def test_same_file_twice_is_one_document():
    store, case = seeded()
    documents = DocumentStore(store)
    first = documents.add(case.id, "x.txt", b"hello")
    assert documents.add(case.id, "y.txt", b"hello").id == first.id


def test_list_of_dates_is_chronological_and_sourced():
    _, case = seeded()
    rows = list_of_dates.rows(case)
    assert [r["date"] for r in rows] == sorted(r["date"] for r in rows)
    assert rows[-1]["event"].startswith("Next date of hearing")
    assert all(r["source"] for r in rows)
    assert list_of_dates.to_docx(case)[:2] == b"PK"


def client():
    store = CaseStore()
    seed(store, today=TODAY)
    return TestClient(create_app(store, ConnectorLadder([]))), store.by_cnr("DLSE010001232024")


def test_api_upload_search_and_export():
    c, case = client()
    uploaded = c.post(
        f"/api/cases/{case.id}/documents",
        files={"file": ("reply.txt", b"Reply of the IO.\fNo previous involvement found.")},
    ).json()
    assert uploaded["pages"] == 2
    matches = c.get(f"/api/cases/{case.id}/search", params={"q": "no previous involvement found"}).json()["matches"]
    assert matches[0]["name"] == "reply.txt" and matches[0]["page"] == 2
    page = c.get(f"/api/documents/{uploaded['id']}/pages/2").json()
    assert page["text"] == "No previous involvement found."
    export = c.get(f"/api/cases/{case.id}/list-of-dates.docx")
    assert export.status_code == 200 and "List of dates" in export.headers["content-disposition"]


def test_ask_without_runtime_is_honest_and_still_useful(monkeypatch):
    monkeypatch.delenv("OPENROUTER_API_KEY", raising=False)
    monkeypatch.delenv("MANU_KIMI_AGENT_BIN", raising=False)
    c, case = client()
    body = c.post(f"/api/cases/{case.id}/ask", json={"question": "recovered"}).json()
    assert body["answer"] is None and body["runtime"] == "unavailable"
    assert body["matches"][0]["name"] == "FIR-512-2025.txt"
    assert c.get("/api/runtime").json()["ready"] is False


def test_questions_search_their_key_words_and_best_pages_win():
    store, case = seeded()
    documents = DocumentStore(store)
    hits = documents.search(case.id, "What was recovered?")
    assert hits and hits[0]["name"] == "FIR-512-2025.txt" and hits[0]["page"] == 2
    assert hits[0]["terms"] == ["recovered"]
    assert documents.search(case.id, "what is the") == []


def test_document_kinds_prefer_the_specific_title():
    store, case = seeded()
    kinds = {d.name: d.kind for d in DocumentStore(store).for_case(case.id)}
    assert kinds == {"Bail-application.txt": "bail_application", "FIR-512-2025.txt": "fir"}
