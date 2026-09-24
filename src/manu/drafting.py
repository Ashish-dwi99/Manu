"""Drafts from templates, filled from the record, paused for what only the advocate knows.

Two templates for v1, the two applications a litigating chamber files most:

* **Adjournment application** — Order XVII Rule 1 CPC (civil) or Section 346 BNSS
  (criminal).
* **Regular bail application** — Section 483 BNSS (Sessions Court / High Court) or
  Section 480 BNSS (Magistrate).

A template never guesses. Facts come from the case record, each paragraph tagged with where
it came from; the Section 479 and custody figures come from the law engine with its
working; judgments come from the authorities the advocate relied on. What the record cannot
say (the reason for the adjournment, the grounds, whether an earlier bail application was
filed) is asked, the way Mike's `ask_inputs` pauses a draft (the idea, not the code).
Until every required answer is in, the draft says what it still needs.

Manu drafts; it never files. Every draft ends "Advocate review required before filing".
"""

from __future__ import annotations

import io
from dataclasses import dataclass, field
from datetime import date
from typing import Literal

from manu import judge
from manu.case_state.models import Case

Kind = Literal["heading", "center", "title", "para", "prayer", "sign", "note"]


@dataclass(slots=True)
class Block:
    kind: Kind
    text: str
    sources: list[str] = field(default_factory=list)

    def as_dict(self) -> dict:
        return {"kind": self.kind, "text": self.text, "sources": self.sources}


@dataclass(frozen=True, slots=True)
class Question:
    key: str
    label: str
    hint: str = ""
    required: bool = True
    kind: Literal["text", "long", "choice"] = "text"
    options: tuple[tuple[str, str], ...] = ()

    def as_dict(self) -> dict:
        return {
            "key": self.key,
            "label": self.label,
            "hint": self.hint,
            "required": self.required,
            "kind": self.kind,
            "options": [{"value": v, "label": label} for v, label in self.options],
        }


@dataclass(slots=True)
class Draft:
    template: str
    title: str
    questions: list[Question]
    blocks: list[Block]
    missing: list[str]
    flags: list[str] = field(default_factory=list)

    @property
    def ready(self) -> bool:
        return not self.missing

    def as_dict(self) -> dict:
        return {
            "template": self.template,
            "title": self.title,
            "status": "ready" if self.ready else "needs_input",
            "questions": [q.as_dict() for q in self.questions],
            "missing": self.missing,
            "blocks": [b.as_dict() for b in self.blocks],
            "flags": self.flags,
        }


def _d(value: date) -> str:
    return value.strftime("%d.%m.%Y")


def _order_src(on: date) -> str:
    return f"order dated {_d(on)}"


def _parties(case: Case) -> list[Block]:
    """The cause title: parties as the record names them."""
    if not case.parties:
        return [Block("center", case.title or case.cnr, ["court record"])]
    first = [p for p in case.parties if p.role in ("petitioner", "appellant", "complainant", "state")]
    second = [p for p in case.parties if p not in first]
    blocks = [Block("center", " & ".join(p.name for p in first) or "—", ["court record"])]
    blocks.append(Block("center", "versus"))
    blocks.append(Block("center", " & ".join(p.name for p in second) or "—", ["court record"]))
    return blocks


def _head(case: Case, heading: str) -> list[Block]:
    court = case.court.upper() if case.court else "IN THE COURT OF ____________"
    if case.court and not court.startswith("IN THE"):
        court = f"IN THE {court}"
    blocks = [Block("heading", court, ["court record"] if case.court else [])]
    if case.case_number:
        blocks.append(Block("center", case.case_number, ["court record"]))
    if case.cnr:
        blocks.append(Block("center", f"CNR No. {case.cnr}", ["court record"]))
    blocks += _parties(case)
    blocks.append(Block("title", heading))
    blocks.append(Block("para", "MOST RESPECTFULLY SHOWETH:"))
    return blocks


def _sign(inputs: dict) -> list[Block]:
    counsel = (inputs.get("counsel") or "").strip() or "____________"
    return [
        Block("sign", f"Through counsel\n{counsel}\nAdvocate", ["you"] if inputs.get("counsel") else []),
        Block("sign", "Place: ____________\nDate: ____________"),
        Block("note", "Drafted with Manu from the court record. Advocate review required before filing."),
    ]


def _authorities(case: Case) -> Block | None:
    if not case.authorities:
        return None
    cites = []
    for a in case.authorities:
        where = f"{a.citation}, para {a.paragraph}" if a.citation else f"para {a.paragraph}"
        cites.append(f"{a.title} [{where}]")
    return Block(
        "para",
        "That the applicant craves leave to rely on " + "; ".join(cites) + ", copies of which will be placed on record.",
        [f"authority: {a.title} ¶{a.paragraph} ({a.standing})" for a in case.authorities],
    )


def _missing(questions: list[Question], inputs: dict) -> list[str]:
    return [q.key for q in questions if q.required and not str(inputs.get(q.key) or "").strip()]


COUNSEL = Question("counsel", "Counsel's name", "As it should appear under the signature", required=False)


# -- adjournment ----------------------------------------------------------------------------


def adjournment(case: Case, inputs: dict, *, as_of: date) -> Draft:
    criminal = case.case_type == "criminal" or bool(case.accused)
    provision = "Section 346 of the Bharatiya Nagarik Suraksha Sanhita, 2023" if criminal else "Order XVII Rule 1 of the Code of Civil Procedure, 1908"
    questions = [
        Question("applicant", "On whose behalf", "e.g. the accused, the defendant, the plaintiff"),
        Question("reason", "Why the adjournment is needed", "Stated plainly: counsel unwell, documents awaited, witness unavailable…", kind="long"),
        Question(
            "previous",
            "Adjournments already taken by this side",
            "The number, for the court's record",
            required=not criminal,
        ),
        COUNSEL,
    ]
    missing = _missing(questions, inputs)
    flags: list[str] = []
    blocks = _head(case, f"APPLICATION FOR ADJOURNMENT UNDER {provision.upper()}")
    applicant = (inputs.get("applicant") or "").strip() or "[on whose behalf]"
    n = 1
    if case.next_date:
        blocks.append(
            Block(
                "para",
                f"{n}. That the above matter is listed before this Hon'ble Court on {_d(case.next_date)}"
                + (f" for {case.next_purpose.lower()}." if case.next_purpose else "."),
                ["court record"],
            )
        )
        n += 1
    order = case.last_order
    if order:
        blocks.append(
            Block(
                "para",
                f"{n}. That on the last date of hearing, {_d(order.on)}, the matter was "
                + (f"taken up and the Court recorded: “{order.title}”." if order.title else "taken up."),
                [_order_src(order.on)],
            )
        )
        n += 1
    reason = (inputs.get("reason") or "").strip() or "[the reason for the adjournment]"
    blocks.append(Block("para", f"{n}. That {reason.rstrip('.')}.", ["you"] if inputs.get("reason") else []))
    n += 1
    previous = str(inputs.get("previous") or "").strip()
    if previous:
        blocks.append(
            Block("para", f"{n}. That the {applicant} has taken {previous} adjournment(s) so far in this matter.", ["you"])
        )
        n += 1
        if not criminal and previous.isdigit() and int(previous) >= 3:
            flags.append(
                "Order XVII Rule 1 proviso: no more than three adjournments to a party during the hearing of a suit. "
                "Say why this one should be allowed."
            )
    blocks.append(
        Block(
            "para",
            f"{n}. That the application is bona fide and made in the interest of justice, and no prejudice will be "
            "caused to the other side, who may be compensated by costs if the Court so directs.",
        )
    )
    blocks.append(
        Block(
            "prayer",
            "PRAYER\nIt is, therefore, most respectfully prayed that this Hon'ble Court may be pleased to adjourn the "
            f"matter to a date convenient to the Court, on behalf of the {applicant}, and pass such other order as it "
            "deems fit in the interest of justice.",
        )
    )
    if criminal:
        flags.append("Section 346 BNSS: in a criminal trial, adjournments are the exception once evidence begins; the court records reasons.")
    blocks += _sign(inputs)
    return Draft("adjournment", "Adjournment application", questions, blocks, missing, flags)


# -- bail -----------------------------------------------------------------------------------


def bail(case: Case, inputs: dict, *, as_of: date) -> Draft:
    assessment = judge.assess_case(case, as_of)
    names = [a.name for a in case.accused]
    questions = []
    if len(names) > 1:
        questions.append(Question("accused", "Applicant", "Which accused is applying", kind="choice", options=tuple((n, n) for n in names)))
    questions += [
        Question(
            "provision",
            "Power invoked",
            "Section 483 BNSS before a Sessions Court or High Court; Section 480 BNSS before a Magistrate",
            kind="choice",
            options=(("483", "Section 483 BNSS (Sessions Court / High Court)"), ("480", "Section 480 BNSS (Magistrate)")),
        ),
        Question(
            "previous_bail",
            "Earlier bail applications",
            "Whether any other bail application was filed or is pending, with court and outcome. Courts require this disclosure.",
            kind="long",
        ),
        Question("grounds", "Further grounds", "Anything the record does not show: role, parity, health, family", required=False, kind="long"),
        Question("roots", "Roots in society", "Residence, family, occupation; willingness to furnish surety", required=False, kind="long"),
        COUNSEL,
    ]
    missing = _missing(questions, inputs)
    flags: list[str] = []
    if assessment is None or not case.accused:
        flags.append("This case has no accused on record, so the bail template has nothing to draw on.")
    chosen = (inputs.get("accused") or (names[0] if len(names) == 1 else "")).strip()
    person = next((a for a in case.accused if a.name == chosen), None)
    graded = next((x for x in (assessment.accused if assessment else []) if x.accused.name == chosen), None)
    provision = {"483": "SECTION 483", "480": "SECTION 480"}.get(str(inputs.get("provision") or ""), "SECTION ___")
    blocks = _head(case, f"APPLICATION FOR REGULAR BAIL UNDER {provision} OF THE BHARATIYA NAGARIK SURAKSHA SANHITA, 2023")
    who = chosen or "[the applicant]"
    n = 1

    def para(text: str, sources: list[str]) -> None:
        nonlocal n
        blocks.append(Block("para", f"{n}. {text}", sources))
        n += 1

    if case.charges:
        para(
            f"That the applicant, {who}, is an accused in the above case, in which the charges on record are "
            + ", ".join(c.text for c in case.charges)
            + ".",
            ["court record"],
        )
    if person and person.custody:
        first = min(span.start for span in person.custody)
        days = graded.s479.days_detained if graded else None
        text = f"That the applicant has been in judicial custody since {_d(first)}"
        text += f", that is, for {days} days as on {_d(as_of)}." if days is not None else "."
        para(text, ["court record (remand dates)"])
    if case.chargesheet_filed_on:
        para(
            f"That the chargesheet was filed on {_d(case.chargesheet_filed_on)}; the investigation is complete and the "
            "applicant is not required for custodial interrogation.",
            ["court record"],
        )
    if person and person.first_offender is True:
        para("That the applicant has no previous involvement in any criminal case.", ["court record (antecedents)"])
    elif person and person.first_offender is None:
        flags.append("The record does not say whether the applicant has previous involvement; the draft is silent on it.")
    if assessment and assessment.offences:
        worst = max(assessment.offences, key=lambda o: (o.punishment.max_is_life_or_death, o.punishment.max_years or 0))
        if not worst.punishment.max_is_life_or_death and worst.punishment.max_years:
            para(
                f"That the maximum punishment for the offences alleged is imprisonment for {worst.punishment.max_years:g} years "
                f"({worst.key}).",
                ["offence table (fixed rules)"],
            )
            if not worst.reviewed:
                flags.append(f"The punishment for {worst.key} is from Manu's offence table, not yet reviewed by an advocate.")
    if graded and graded.s479.applicable and graded.s479.status in ("approaching", "crossed", "maximum_exceeded"):
        t = graded.s479.applicable
        verb = "has undergone" if graded.s479.status != "approaching" else "will have undergone"
        para(
            f"That the applicant {verb} {t.ratio_label} of the maximum sentence on {_d(t.crossing_date)}, and is entitled "
            "to be considered for release under Section 479 BNSS.",
            ["Section 479 working (fixed rules)"],
        )
    grounds = (inputs.get("grounds") or "").strip()
    if grounds:
        para(f"That {grounds.rstrip('.')}.", ["you"])
    roots = (inputs.get("roots") or "").strip()
    if roots:
        para(f"That {roots.rstrip('.')}. The applicant undertakes to abide by any condition the Court imposes.", ["you"])
    else:
        para("That the applicant undertakes to abide by any condition this Hon'ble Court may impose and to furnish surety.", [])
    earlier = (inputs.get("previous_bail") or "").strip()
    para(
        f"That {earlier.rstrip('.')}." if earlier else "That [earlier bail applications: filed / pending / none].",
        ["you"] if earlier else [],
    )
    authority = _authorities(case)
    if authority:
        authority.text = f"{n}. {authority.text}"
        blocks.append(authority)
        n += 1
        if any(a.standing == "demo" for a in case.authorities):
            flags.append("An authority cited is from the fictional demo library. Remove it before filing.")
        if any(a.standing == "lead" for a in case.authorities):
            flags.append("An authority cited is a lead: confirm it from an official copy before filing.")
    blocks.append(
        Block(
            "prayer",
            "PRAYER\nIt is, therefore, most respectfully prayed that this Hon'ble Court may be pleased to release the "
            f"applicant, {who}, on regular bail in the above case on such terms and conditions as it deems fit, in the "
            "interest of justice.",
        )
    )
    blocks += _sign(inputs)
    return Draft("bail", "Regular bail application", questions, blocks, missing, flags)


TEMPLATES = {
    "adjournment": ("Adjournment application", adjournment, lambda case: True),
    "bail": ("Regular bail application", bail, lambda case: bool(case.accused)),
}


def available(case: Case) -> list[dict]:
    return [{"key": key, "title": title} for key, (title, _build, applies) in TEMPLATES.items() if applies(case)]


def build(template: str, case: Case, inputs: dict, *, as_of: date) -> Draft:
    if template not in TEMPLATES:
        raise KeyError(template)
    _title, builder, applies = TEMPLATES[template]
    if not applies(case):
        raise ValueError("This template does not apply to this case.")
    return builder(case, {k: v for k, v in (inputs or {}).items() if v is not None}, as_of=as_of)


def to_docx(draft: Draft) -> bytes:
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Pt

    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Times New Roman"
    style.font.size = Pt(12)
    for block in draft.blocks:
        for i, line in enumerate(block.text.split("\n")):
            para = doc.add_paragraph()
            run = para.add_run(line)
            if block.kind in ("heading", "center", "title"):
                para.alignment = WD_ALIGN_PARAGRAPH.CENTER
                run.bold = block.kind != "center" or line == "versus"
                run.underline = block.kind == "title"
            elif block.kind == "prayer" and i == 0:
                para.alignment = WD_ALIGN_PARAGRAPH.CENTER
                run.bold = True
            elif block.kind == "sign":
                para.alignment = WD_ALIGN_PARAGRAPH.RIGHT if "counsel" in block.text.lower() else WD_ALIGN_PARAGRAPH.LEFT
            elif block.kind == "note":
                run.italic = True
                run.font.size = Pt(9)
            else:
                para.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    buffer = io.BytesIO()
    doc.save(buffer)
    return buffer.getvalue()
