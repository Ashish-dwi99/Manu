"""The list of dates and events, built from the record, exported as a Word file.

Every Indian petition, appeal and written submission opens with one. Chambers build it by
hand from the file; Manu builds it from what it already holds: offence date, custody,
chargesheet, every hearing and order, the next date. Each row names its source, so the
advocate can check it before it goes into a filing. No model is involved.
"""

from __future__ import annotations

import io
from datetime import date

from manu.case_state.models import Case


def rows(case: Case) -> list[dict]:
    items: list[tuple[date, str, str]] = []
    offence_dates = sorted({c.offence_date for c in case.charges if c.offence_date})
    if offence_dates:
        charges = ", ".join(c.text for c in case.charges)
        items.append((offence_dates[0], f"Date of alleged offence ({charges}).", "Charges on record"))
    for person in case.accused:
        for span in person.custody:
            items.append(
                (
                    span.start,
                    f"{person.name} taken into custody{f' ({span.place})' if span.place else ''}.",
                    "Case record",
                )
            )
            if span.end:
                items.append((span.end, f"{person.name} released from custody.", "Case record"))
    if case.chargesheet_filed_on:
        items.append((case.chargesheet_filed_on, "Chargesheet filed.", "Case record"))
    orders_by_date = {o.on: o for o in case.orders}
    for hearing in case.hearings:
        if hearing.on in orders_by_date:
            continue
        text = hearing.purpose + (f": {hearing.outcome}" if hearing.outcome else "")
        items.append((hearing.on, f"Hearing — {text}.", "Court record"))
    for order in case.orders:
        what = order.title or "Order passed"
        if order.next_date:
            what += f"; matter listed on {order.next_date.strftime('%d.%m.%Y')}" + (
                f" for {order.next_purpose}" if order.next_purpose else ""
            )
        items.append((order.on, f"{what}.", f"Order dated {order.on.strftime('%d.%m.%Y')}"))
    if case.next_date:
        items.append(
            (case.next_date, f"Next date of hearing — {case.next_purpose or 'purpose not recorded'}.", "Court record")
        )
    items.sort(key=lambda item: item[0])
    return [{"date": d.isoformat(), "display": d.strftime("%d.%m.%Y"), "event": e, "source": s} for d, e, s in items]


def to_docx(case: Case) -> bytes:
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Pt

    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Times New Roman"
    style.font.size = Pt(12)

    for line in filter(None, [case.court.upper(), case.case_number, case.title]):
        para = doc.add_paragraph(line)
        para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        para.runs[0].bold = True
    heading = doc.add_paragraph("LIST OF DATES AND EVENTS")
    heading.alignment = WD_ALIGN_PARAGRAPH.CENTER
    heading.runs[0].bold = True
    heading.runs[0].underline = True

    table = doc.add_table(rows=1, cols=3)
    table.style = "Table Grid"
    for cell, text in zip(table.rows[0].cells, ("Date", "Event", "Source"), strict=True):
        cell.text = text
        cell.paragraphs[0].runs[0].bold = True
    for row in rows(case):
        cells = table.add_row().cells
        cells[0].text, cells[1].text, cells[2].text = row["display"], row["event"], row["source"]

    note = doc.add_paragraph("Prepared by Manu from the court record. Advocate review required before filing.")
    note.runs[0].italic = True
    note.runs[0].font.size = Pt(9)
    buffer = io.BytesIO()
    doc.save(buffer)
    return buffer.getvalue()
